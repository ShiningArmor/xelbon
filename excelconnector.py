from openpyxl import load_workbook
from utils import xls2xlsx, index2excel
import json

datatypes = {
    "int": int,
    "float": float,
    "unicode": unicode,
    "str": str,
    "None": None,
}

form_datatypes = {
    "int": "Entero",
    "float": "Decimal",
    "unicode": "Texto",
    "str": "Texto",
    "None": "Vacio",
}


class ExcelConnector(object):
    def __init__(self, pathfile, mode="r"):
        self.pathfile = pathfile
        self.open_mode = mode
        self.table_index = 0
        self.wb = None
        self.tables = []
        self.forms = []
        self.open()
        self.initialize()

    def open(self):
        readonly = True if self.open_mode == "r" else False
        if self.pathfile[self.pathfile.rfind("."):] == ".xls":
            self.wb = xls2xlsx(self.pathfile)
        elif self.pathfile[self.pathfile.rfind("."):] == ".xlsx":
            self.wb = load_workbook(filename=self.pathfile, read_only=readonly)

    def initialize(self):
        sheets_names = self.wb.get_sheet_names()
        for idx, sheet in enumerate(sheets_names):
            if "FormWeb" in sheet:
                table = Table(self.get_sheet(idx), idx)
                self.forms.append(table)
            else:
                table = Table(self.get_sheet(idx),idx)
                self.tables.append(table)

    def set_table_crusors(self):
        for table in self.tables:
            table.set_cursor()

    def get_sheet(self, i):
        return  self.wb[self.wb.get_sheet_names()[i]]

class Table(object):
    def __init__(self, sheet_obj, index):
        self.sheet = sheet_obj
        self.name = self.sheet.title
        self.wb_index = index
        self.dimensions = self.sheet.dimensions
        self.indexes = self.find_table()
        self.cursor = [0, 0]

    def set_cursor(self):
        self.cursor = [self.indexes[0]["column"], len(self.sheet[self.indexes[0]["column"]]) + 1]

    def refresh(self):
        pass

    def get_field_dataype(self, cell):
        for _type in [int,float,unicode,str, None]:
            if isinstance(self.sheet[index2excel(cell.row  + 1, cell.col_idx)].value, _type):
                a = str(_type)
                if _type is not None:
                    a = a[a.find("'") + 1:a.rfind("'")]
                return a

        return 'None'

    def find_table(self):
        stop = False
        indexes = []
        for row in self.sheet.rows:
            for i in row:
                if i.value:
                    indexes.append({"title": i.value,
                                    "column": i.column,
                                    "row": i.row,
                                    "datatype": self.get_field_dataype(i)
                                   })
                    stop = True
            if stop:
                break
        return indexes

    def export(self):
        return json.dumps(dict(
            name= self.name,
            filexls=self.sheet.title,
            cursor=self.cursor,
            fields=self.indexes
        ))


class Form(object):
    def __init__(self, sheet):
        self.sheet = sheet
        self.name = self.sheet.title
        self.dimensions = self.sheet.dimensions
        self.title = ""


    def get_form(self):
        col = 1
        fil = 2



    def create_form(self):
        form = [["FORMULARIO WEB: %s" % self.name,],]
        for index in self.indexes:
            row = []
            row.append(index["title"])
            row.append(form_datatypes[index["datatype"]])
            form.append(row)
        return form